import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from decimal import Decimal
from unittest import mock

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
sys.path.insert(0, HERE)

from docpipe import normalize as N  # noqa: E402
from docpipe import ooxml, pdf, pipeline, report, sniff, store  # noqa: E402
import fixtures as F  # noqa: E402

SALES = {
    "مبيعات": [
        ["الفرع", "التاريخ", "المبلغ", "ملاحظة"],
        ["القاهرة", ("n", "46325"), 1250.5, "نقدي"],
        ["الإسكندرية", ("n", "46326"), "١٬٢٠٠٫٢٥", "تحويل"],
        ["الجيزة", ("n", "46327"), "(300)", "مرتجع"],
        ["المجموع", None, ("n", "2150.75"), None],
    ],
}


def tmp():
    d = tempfile.mkdtemp(prefix="docpipe-")
    return d


class Normalize(unittest.TestCase):
    def test_numbers(self):
        cases = {"1,234.50": "1234.50", "(1,234)": "-1234", "١٢٣٫٥": "123.5", "12%": "0.12", "EGP 1,000": "1000",
                 "٣٠٠ ج.م": "300", "12,5": "12.5", "-7": "-7", "١٬٢٠٠٫٢٥": "1200.25"}
        for raw, want in cases.items():
            v, _ = N.parse_number(raw)
            self.assertEqual(v, Decimal(want), raw)
        for raw in ("abc", "", None, "12/10/2026", "عميل 17", "Invoice 1001", "Q3", "فرع 5"):
            self.assertIsNone(N.parse_number(raw)[0], raw)

    def test_money_is_exact(self):
        total = sum(Decimal("0.1") for _ in range(10))
        self.assertEqual(total, Decimal("1.0"))
        self.assertEqual(N.money("2.675"), Decimal("2.68"))  # half-up, a float would give 2.67

    def test_dates(self):
        self.assertEqual(N.parse_date("30/10/2026")[:2], ("2026-10-30", 1.0))
        self.assertEqual(N.parse_date("03/04/2026")[:2], ("2026-04-03", 0.6))  # ambiguous: lower confidence
        self.assertEqual(N.parse_date("١٥/٠٩/٢٠٢٦")[0], "2026-09-15")
        self.assertEqual(N.excel_serial_to_date(46325).date().isoformat(), "2026-10-30")


class Readers(unittest.TestCase):
    def setUp(self):
        self.d = tmp()

    def tearDown(self):
        shutil.rmtree(self.d)

    def test_xlsx_values_provenance_flags(self):
        p = os.path.join(self.d, "sales.xlsx")
        F.make_xlsx(p, SALES, hidden_rows={("مبيعات", 4)}, date_cells={("مبيعات", f"B{r}") for r in (2, 3, 4)},
                    formulas={("مبيعات", "C5"): "SUM(C2:C4)"}, merged={"مبيعات": ["A5:B5"]})
        recs = {r["cell"]: r for r in ooxml.xlsx_records(p)}
        self.assertEqual(recs["A2"]["raw"], "القاهرة")
        self.assertTrue(recs["B2"]["is_date"])
        self.assertTrue(recs["A4"]["row_hidden"])
        self.assertTrue(recs["C5"]["formula"])
        self.assertEqual(recs["A5"]["merged"], "A5:B5")
        try:
            import openpyxl  # independent oracle, not used by docpipe
        except ImportError:
            self.skipTest("openpyxl not installed (oracle only)")
        ws = openpyxl.load_workbook(p, data_only=True)["مبيعات"]
        for ref, r in recs.items():
            o = ws[ref].value
            if hasattr(o, "date"):
                self.assertEqual(N.excel_serial_to_date(r["raw"]).date(), o.date(), ref)
            elif isinstance(o, (int, float)):
                self.assertEqual(Decimal(r["raw"]), Decimal(str(o)), ref)
            else:
                self.assertEqual(r["raw"], o, ref)

    def test_docx_and_pptx_against_oracles(self):
        dp = os.path.join(self.d, "memo.docx")
        F.make_docx(dp, ["تقرير الربع الثالث", "إجمالي المبيعات 2,150.75 جنيه"], table=[["البند", "القيمة"], ["إيجار", "5,000"]], rtl=True)
        recs = list(ooxml.docx_records(dp))
        self.assertEqual(recs[0]["where"], "p1")
        self.assertTrue(recs[0]["rtl"])
        self.assertIn({"where": "t1r2c2", "text": "5,000", "rtl": False, "style": "table"}, recs)
        pp = os.path.join(self.d, "deck.pptx")
        F.make_pptx(pp, [("الملخص", ["نمو 12%", "هامش 30%"], None), ("الأرقام", [], [["ربع", "إيراد"], ["Q3", "2,150.75"]])])
        prec = list(ooxml.pptx_records(pp))
        self.assertIn({"where": "s2:Table 3:r2c2", "text": "2,150.75"}, prec)
        try:
            import docx
            import pptx
        except ImportError:
            self.skipTest("python-docx/pptx not installed (oracles only)")
        self.assertEqual([r["text"] for r in recs if r["where"].startswith("p")], [p.text for p in docx.Document(dp).paragraphs if p.text.strip()])
        texts = []
        for s in pptx.Presentation(pp).slides:
            for sh in s.shapes:
                if sh.has_text_frame:
                    texts += [p.text for p in sh.text_frame.paragraphs if p.text.strip()]
        self.assertEqual(sorted(texts), sorted(r["text"] for r in prec if ":p" in r["where"]))

    def test_pdf_text_and_scanned_page(self):
        p = os.path.join(self.d, "inv.pdf")
        F.make_pdf(p, [["Invoice 1001", "Total EGP 2,150.75"], ["Page two (flate)", "Paid"]])
        recs = list(pdf.pdf_records(p))
        self.assertEqual([(r["page"], r["text"]) for r in recs], [(1, "Invoice 1001"), (1, "Total EGP 2,150.75"), (2, "Page two (flate)"), (2, "Paid")])
        if shutil.which("pdftotext"):
            out = subprocess.run(["pdftotext", "-layout", p, "-"], capture_output=True, text=True).stdout
            for r in recs:
                self.assertIn(r["text"], out)
        blank = os.path.join(self.d, "scan.pdf")
        F.make_pdf(blank, [[]])
        r = list(pdf.pdf_records(blank))
        self.assertEqual(r[0]["status"], "needs-reader")  # no text layer: reported, never invented

    def test_libreoffice_made_pdf(self):
        if not shutil.which("soffice"):
            self.skipTest("LibreOffice not installed")
        dp = os.path.join(self.d, "lo.docx")
        F.make_docx(dp, ["Quarterly total 2150.75", "Second line", "إجمالي المبيعات"])
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", self.d, dp], capture_output=True, timeout=120)
        out = os.path.join(self.d, "lo.pdf")
        if not os.path.exists(out):
            self.skipTest("conversion failed")
        recs = list(pdf.pdf_records(out))
        text = " ".join(r.get("text", "") for r in recs)
        # LibreOffice embeds subset fonts with a ToUnicode map: read through it, Arabic in reading order.
        self.assertIn("Quarterly total 2150.75", text)
        self.assertIn("إجمالي المبيعات", text)


class Pipeline(unittest.TestCase):
    def setUp(self):
        self.d = tmp()
        self.inp = os.path.join(self.d, "in")
        os.makedirs(self.inp)
        F.make_xlsx(os.path.join(self.inp, "sales.xlsx"), SALES, date_cells={("مبيعات", f"B{r}") for r in (2, 3, 4)},
                    formulas={("مبيعات", "C5"): "SUM(C2:C4)"})
        self.db = store.connect(os.path.join(self.d, "w.db"))

    def tearDown(self):
        self.db.close()
        shutil.rmtree(self.d)

    def test_totals_provenance_and_idempotent_import(self):
        p = os.path.join(self.inp, "sales.xlsx")
        _, st = pipeline.import_file(self.db, p)
        self.assertEqual(st, "done")
        t = pipeline.totals(self.db, "المبلغ")["-"]
        self.assertEqual(t["total"], Decimal("1250.5") + Decimal("1200.25") - Decimal("300") + Decimal("2150.75"))
        self.assertIn((1, "مبيعات!C3"), t["sources"])  # every number traces to its cell
        flags = self.db.execute("SELECT flags FROM cells WHERE loc='مبيعات!C5'").fetchone()[0]
        self.assertIn("not recalculated", flags)
        n = self.db.execute("SELECT COUNT(*) FROM cells").fetchone()[0]
        _, st2 = pipeline.import_file(self.db, p)
        self.assertEqual(st2, "already-imported")
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM cells").fetchone()[0], n)

    def test_cut_import_resumes_without_duplicates(self):
        big = os.path.join(self.inp, "big.xlsx")
        rows = (["فرع %d" % i, i, f"{i}.25"] for i in range(1, 5001))
        F.make_xlsx(big, {"data": [["name", "n", "amount"]]}, stream_rows=("data", rows))
        with self.assertRaises(pipeline.Crash):
            pipeline.import_file(self.db, big, batch=1000, crash_after_batches=3)
        cut = self.db.execute("SELECT COUNT(*) FROM cells").fetchone()[0]
        self.assertGreater(cut, 0)
        _, st = pipeline.import_file(self.db, big, batch=1000)
        self.assertEqual(st, "done")
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM cells WHERE file_id=(SELECT id FROM files WHERE path=?)", (big,)).fetchone()[0], 3 + 5000 * 3)
        t = pipeline.totals(self.db, "amount")["-"]
        self.assertEqual(t["total"], sum(Decimal(f"{i}.25") for i in range(1, 5001)))
        self.assertEqual(t["count"], 5000)

    def test_company_profile_never_parses_office_bytes(self):
        p = os.path.join(self.inp, "sales.xlsx")
        with mock.patch.object(ooxml, "xlsx_records", side_effect=AssertionError("parsed a protected file")), \
             mock.patch.object(pipeline, "com_extract", return_value=(None, "no COM here")) as com:
            _, st = pipeline.import_file(self.db, p, profile="company_nasca")
        self.assertEqual(st, "needs-authorized-app")
        com.assert_called_once()
        row = self.db.execute("SELECT route, status, note FROM files").fetchone()
        self.assertEqual(row[:2], ("com", "needs-authorized-app"))

    def test_company_profile_uses_com_records(self):
        p = os.path.join(self.inp, "sales.xlsx")
        fake = [{"loc": "S!A1", "sheet": "S", "row": 1, "col": 1, "raw": "amount", "kind": "s"},
                {"loc": "S!A2", "sheet": "S", "row": 2, "col": 1, "raw": 10.5, "kind": "n"},
                {"loc": "S!A3", "sheet": "S", "row": 3, "col": 1, "raw": 4.5, "kind": "n", "formula": True}]
        with mock.patch.object(pipeline, "com_extract", return_value=(fake, None)):
            _, st = pipeline.import_file(self.db, p, profile="company_nasca")
        self.assertEqual(st, "done")
        self.assertEqual(pipeline.totals(self.db, "amount")["-"]["total"], Decimal("15.0"))
        self.assertEqual(self.db.execute("SELECT method FROM cells WHERE loc='S!A2'").fetchone()[0], "office-com")

    def test_zip_bomb_and_cfb_are_not_parsed(self):
        bomb = os.path.join(self.inp, "bomb.xlsx")
        import zipfile
        with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("xl/workbook.xml", "<workbook/>")
            z.writestr("xl/worksheets/sheet1.xml", "0" * (60 << 20))
        _, st = pipeline.import_file(self.db, bomb)
        self.assertEqual(st, "quarantined")
        enc = os.path.join(self.inp, "protected.xlsx")
        with open(enc, "wb") as f:
            f.write(sniff.CFB + b"\0" * 4096)  # an encrypted/DRM OOXML is a CFB container
        with mock.patch.object(pipeline, "com_extract", return_value=(None, "no COM here")):
            _, st = pipeline.import_file(self.db, enc)
        self.assertEqual(st, "needs-authorized-app")  # not "damaged", not parsed

    def test_report_and_reference_check(self):
        pipeline.import_file(self.db, os.path.join(self.inp, "sales.xlsx"))
        html_p, csv_p = os.path.join(self.d, "r.html"), os.path.join(self.d, "r.csv")
        res = report.build(self.db, html_p, csv_p, "تقرير", ["المبلغ"], {"المبلغ": "4301.50", "x": "1"})
        self.assertEqual(res["checks"], [("المبلغ", "4301.50", "4301.50", True)])
        bad = report.build(self.db, html_p, csv_p, "تقرير", ["المبلغ"], {"المبلغ": "4301.49"})
        self.assertFalse(bad["checks"][0][3])
        with open(html_p, encoding="utf-8") as fh:
            body = fh.read()
        self.assertIn('dir="rtl"', body)
        self.assertIn("مبيعات!C2", body)  # provenance shown
        import csv
        with open(csv_p, encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.reader(fh))
        self.assertEqual(rows[0][0], "file")
        self.assertTrue(any(r[3] == "المبلغ" and r[4] == "1200.25" and r[8] == "مبيعات!C3" for r in rows))

    def test_backup_and_integrity(self):
        pipeline.import_file(self.db, os.path.join(self.inp, "sales.xlsx"))
        bk = os.path.join(self.d, "b.db")
        store.backup(self.db, bk)
        b = store.connect(bk)
        try:
            self.assertEqual(store.integrity(b), "ok")
            self.assertEqual(b.execute("SELECT COUNT(*) FROM cells").fetchone(), self.db.execute("SELECT COUNT(*) FROM cells").fetchone())
            self.assertEqual(b.execute("PRAGMA user_version").fetchone()[0], len(store.MIGRATIONS))
        finally:
            b.close()  # Windows cannot delete an open database file


if __name__ == "__main__":
    unittest.main()
